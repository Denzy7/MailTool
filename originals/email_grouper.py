#!/usr/bin/env python3
"""
Email Grouper
=============

Reads a CSV of emails exported in the format:

    Sender, Email, Date Received, Subject, Body

and assigns each row to a "group" by keyword matching.

Matching order
--------------
1. Subject + Body are searched for each group's keywords and, as a fallback
   (or first, if the toggle is flipped), the group's literal name.
2. If nothing matched, the original message is located via the chosen mail
   source - disambiguated by exact timestamp when several candidates share
   the same subject/sender - and its pdf / docx / doc attachments are
   downloaded/read, text is extracted and searched the same way.
3. If still nothing matched, the row is marked UNMATCHED.

Mail source (Run tab)
----------------------
Chosen per run, with a toggle:
  IMAP        connects to the server configured on the IMAP Login tab.
  Local export folder   scans a folder of .eml files recursively - e.g. an
              extracted Zimbra tgz export (Inbox/, Inbox!1/, Inbox!2/, ...;
              .meta sidecars are ignored, only *.eml is read). No password
              needed. Point it at the extraction root; subfolders are found
              automatically.
Both sources share the same IMAP cache (imap_cache.json), exclude/include
attachment filters, Reply-To handling, and exact-timestamp disambiguation.

Keyword layout (Groups tab)
---------------------------
Each group has a name plus a list of keyword lines:
    line 1   the descriptive name, written to the CSVs as "Descriptive Name".
             It is also matched against the CSV subject + body AND attachments.
    line 2   comma-separated keywords matched against IMAP attachment text
             (live download and IMAP cache) AND the CSV subject + body.
             e.g.  jan, january, jaanuary, njan,
    line 3+  one keyword per line, matched against the CSV subject + body only.
The group name itself is a fallback match in both places.

Match precedence is a toggle on the Groups tab: "keywords first, then group
name" (default) or "group name first, then keywords".

Reply-To
--------
When a message is fetched (IMAP or local .eml) for attachment search, its
Reply-To header (if present) is used as the row's Name/Email instead of the
CSV's Sender/Email columns.

Attachment filters
------------------
Two pattern lists on the "Attachment Filters" tab: EXCLUDE skips a matching
attachment entirely (never downloaded/opened/searched); INCLUDE overrides
EXCLUDE for anything matching both, forcing it to download anyway - e.g.
excluding "notes" would normally drop "notes_reports_and_work.pdf", but
including that name (or a pattern like "notes_reports*") forces it through.

Output
------
matched.csv   ->  "Name, Email, Descriptive Name, Group Name, Date, Subject"
unmatched.csv ->  "Name, Email, Subject, Date Received, Reason"

Config
------
Groups + IMAP host/port/username/mailbox are persisted to config.json.
The password is NEVER written to disk - it is asked for at run time and kept
in memory only.

Logging
-------
Every run writes a timestamped log file to <output folder>/logs/
(run_YYYY-MM-DD_HHMMSS.log), mirroring everything shown in the on-screen log.

Cache
-----
IMAP lookups (message found / not found, and every attachment's extracted
text) are cached to imap_cache.json, keyed by sender email + normalized
subject + date. Reruns against the same CSV reuse this cache instead of
hitting the mailbox again. Use the "Clear cache" button on the Run tab to
force fresh lookups.

Optional third-party libs (install what you need):
    pip install pypdf python-docx docx2txt olefile python-dateutil
For legacy .doc, `antiword` (apt install antiword) gives the best results.
"""

import csv
import email
import email.parser
import email.policy
import imaplib
import json
import os
import fnmatch
import re
import shutil
import subprocess
import tempfile
import threading
import queue
import traceback
from datetime import datetime, timedelta, timezone
from email.header import decode_header, make_header

import tkinter as tk
from tkinter import ttk, filedialog, messagebox

# --------------------------------------------------------------------------
# Optional dependencies - the app degrades gracefully if they are missing
# --------------------------------------------------------------------------
try:
    from pypdf import PdfReader
    HAVE_PDF = True
except ImportError:
    try:
        from PyPDF2 import PdfReader  # type: ignore
        HAVE_PDF = True
    except ImportError:
        HAVE_PDF = False

try:
    import docx  # python-docx
    HAVE_DOCX = True
except ImportError:
    HAVE_DOCX = False

try:
    import docx2txt
    HAVE_DOCX2TXT = True
except ImportError:
    HAVE_DOCX2TXT = False

try:
    import olefile
    HAVE_OLEFILE = True
except ImportError:
    HAVE_OLEFILE = False

try:
    from dateutil import parser as dateparser
    HAVE_DATEUTIL = True
except ImportError:
    HAVE_DATEUTIL = False


APP_DIR = os.path.dirname(os.path.abspath(__file__))
CONFIG_PATH = os.path.join(APP_DIR, "config.json")
CACHE_PATH = os.path.join(APP_DIR, "imap_cache.json")

ALLOWED_EXTS = (".pdf", ".docx", ".doc")

DEFAULT_CONFIG = {
    "imap": {
        "server": "",
        "port": 993,
        "username": "",
        "use_ssl": True,
        "mailbox": "INBOX",
    },
    "local": {
        "folder": "",   # extracted mail export root (e.g. a Zimbra tgz), scanned for .eml files
    },
    "recent": {
        "last_csv": "",
        "last_out_dir": "",
    },
    "options": {
        "whole_word": True,
        "case_sensitive": False,
        "fetch_attachments": True,
        "offline_mode": False,
        "use_cache": True,
        "date_window_days": 1,
        "timestamp_window_seconds": 4,
        "match_precedence": "keywords_first",
        "dedupe_by_email": False,
        "email_source": "imap",   # "imap" or "local"
    },
    "groups": [
        {"group_name": "G1/1/1/2026/01", "keywords": ["G1: January", "jan, january"]},
        {"group_name": "G2/2/1/2026/02", "keywords": ["G2: February", "feb, february"]},
    ],
    "excluded_attachments": [],
    "included_attachments": [],
}


# ==========================================================================
# Config
# ==========================================================================
def load_config():
    if not os.path.exists(CONFIG_PATH):
        return json.loads(json.dumps(DEFAULT_CONFIG))
    try:
        with open(CONFIG_PATH, "r", encoding="utf-8") as fh:
            cfg = json.load(fh)
    except Exception:
        return json.loads(json.dumps(DEFAULT_CONFIG))

    merged = json.loads(json.dumps(DEFAULT_CONFIG))
    for section in ("imap", "local", "options", "recent"):
        if isinstance(cfg.get(section), dict):
            merged[section].update(cfg[section])
    if isinstance(cfg.get("groups"), list):
        merged["groups"] = cfg["groups"]
    if isinstance(cfg.get("excluded_attachments"), list):
        merged["excluded_attachments"] = cfg["excluded_attachments"]
    if isinstance(cfg.get("included_attachments"), list):
        merged["included_attachments"] = cfg["included_attachments"]
    # Defensive: never keep a password even if one was hand-added to the file
    merged["imap"].pop("password", None)
    return merged


def save_config(cfg):
    safe = json.loads(json.dumps(cfg))
    safe.get("imap", {}).pop("password", None)
    with open(CONFIG_PATH, "w", encoding="utf-8") as fh:
        json.dump(safe, fh, indent=2)


# ==========================================================================
# Attachment / lookup cache
# ==========================================================================
# Keyed by (sender email, normalized subject, date string) so repeat runs
# against the same CSV never have to hit the IMAP server again for a
# message already resolved once. Caches both "found attachments" and
# "message not found" so both outcomes are reused.
def load_cache():
    if not os.path.exists(CACHE_PATH):
        return {}
    try:
        with open(CACHE_PATH, "r", encoding="utf-8") as fh:
            data = json.load(fh)
        return data if isinstance(data, dict) else {}
    except Exception:
        return {}


def save_cache(cache):
    tmp = CACHE_PATH + ".tmp"
    with open(tmp, "w", encoding="utf-8") as fh:
        json.dump(cache, fh, indent=2)
    os.replace(tmp, CACHE_PATH)


def cache_key(email_addr, subject, date_s):
    # Version the key so cache entries created by the old "first UID wins"
    # lookup can never be reused by the exact-timestamp matcher.
    subj = re.sub(r"^(re|fwd|fw)\s*:\s*", "", (subject or "").strip(), flags=re.I).strip().lower()
    return f"v2|{(email_addr or '').strip().lower()}|{subj}|{(date_s or '').strip()}"


def normalize_dt(dt):
    """Make an aware/naive datetime comparable with others: convert aware
    datetimes to UTC and strip tzinfo; leave naive ones as-is."""
    if dt is None:
        return None
    if dt.tzinfo is not None:
        return dt.astimezone(timezone.utc).replace(tzinfo=None)
    return dt


def datetime_delta_seconds(a, b):
    """Return absolute seconds between two datetimes."""
    if a is None or b is None:
        return None
    if a.tzinfo is not None and b.tzinfo is not None:
        a = a.astimezone(timezone.utc)
        b = b.astimezone(timezone.utc)
    else:
        a = a.replace(tzinfo=None)
        b = b.replace(tzinfo=None)
    return abs((a - b).total_seconds())


# ==========================================================================
# Keyword matching
# ==========================================================================
def build_pattern(keyword, whole_word=True, case_sensitive=False):
    kw = keyword.strip()
    if not kw:
        return None
    escaped = re.escape(kw)
    if whole_word:
        # \b only works next to word chars; guard both ends.
        prefix = r"\b" if re.match(r"\w", kw[0]) else ""
        suffix = r"\b" if re.match(r"\w", kw[-1]) else ""
        escaped = prefix + escaped + suffix
    flags = 0 if case_sensitive else re.IGNORECASE
    return re.compile(escaped, flags)


def split_keyword_lines(lines):
    """Split a group's raw keyword lines into
    (descriptive_name, csv_keywords, extra_keywords).

    Line 1   the descriptive name (written to the CSVs). Also matched against
             IMAP attachment text AND the CSV subject/body.
    Line 2   comma-separated keywords (trailing commas / blanks are ignored),
             matched against IMAP attachment text (live download + IMAP cache)
             AND the CSV subject/body.
    Line 3+  ONE keyword per line (commas are kept as-is), matched against the
             CSV subject/body only."""
    lines = [l.strip() for l in (lines or []) if l and l.strip()]
    if not lines:
        return "", [], []
    descriptive = lines[0]
    csv_kws = ([k.strip() for k in lines[1].split(",") if k.strip()]
               if len(lines) > 1 else [])
    extra_kws = lines[2:]
    return descriptive, csv_kws, extra_kws


class Group:
    def __init__(self, group_name, keywords):
        self.group_name = group_name
        self.raw_keywords = [k for k in keywords if k.strip()]   # as typed
        self.descriptive, self.csv_keywords, self.extra_keywords = \
            split_keyword_lines(self.raw_keywords)
        self.imap_patterns = []      # line 1 + 2      -> IMAP attachment text / cache
        self.body_patterns = []      # line 1 + 2 + 3+ -> CSV subject + body
        self.group_name_pattern = None

    @property
    def descriptive_name(self):
        """Line 1 of the keyword box."""
        return self.descriptive

    @staticmethod
    def _compile_terms(terms, whole_word, case_sensitive):
        out, seen = [], set()
        for kw in terms:
            key = kw if case_sensitive else kw.lower()
            if not kw or key in seen:
                continue
            seen.add(key)
            pat = build_pattern(kw, whole_word, case_sensitive)
            if pat:
                out.append((kw, pat))
        return out

    def compile(self, whole_word, case_sensitive):
        self.group_name_pattern = build_pattern(self.group_name, whole_word, case_sensitive)
        # The descriptive name (line 1) is searched too, ahead of the other keywords.
        descriptive = [self.descriptive] if self.descriptive else []
        # Lines 1 + 2 -> attachments / IMAP cache.
        self.imap_patterns = self._compile_terms(
            descriptive + self.csv_keywords, whole_word, case_sensitive)
        # Lines 1 + 2 + 3+ -> subject / body.
        self.body_patterns = self._compile_terms(
            descriptive + self.csv_keywords + self.extra_keywords, whole_word, case_sensitive)


def match_by_keywords(text, groups, scope="body"):
    """Check every group's keywords (never the group name).
    scope='body' -> lines 1-3+ (subject/body);  scope='imap' -> lines 1-2 (attachments/cache)."""
    if not text:
        return None, None
    attr = "imap_patterns" if scope == "imap" else "body_patterns"
    for g in groups:
        for kw, pat in getattr(g, attr):
            if pat.search(text):
                return g, kw
    return None, None


def match_by_group_name(text, groups):
    """Check only the literal group name (e.g. 'G1/1/1/2026/01'), never keywords."""
    if not text:
        return None, None
    for g in groups:
        if g.group_name_pattern and g.group_name_pattern.search(text):
            return g, g.group_name
    return None, None


def match_with_precedence(text, groups, precedence="keywords_first", scope="body"):
    """Try one matcher first, fall back to the other. Returns
    (group, matched_term, stage) where stage is 'keyword' / 'imap keyword' /
    'group name', or None when nothing matched."""
    kw_stage = "imap keyword" if scope == "imap" else "keyword"

    def by_keywords(t, gs):
        return match_by_keywords(t, gs, scope)

    if precedence == "groupname_first":
        order = [(match_by_group_name, "group name"), (by_keywords, kw_stage)]
    else:
        order = [(by_keywords, kw_stage), (match_by_group_name, "group name")]
    for matcher, stage in order:
        g, term = matcher(text, groups)
        if g:
            return g, term, stage
    return None, None, None


def whole_word_near_misses(text, groups, case_sensitive=False, limit=5):
    """Attachment keywords that DO appear in `text`, but only inside a longer
    word (e.g. 'jaan' inside 'jaan2026' or 'x_jaan'). Call this after a normal
    match failed: with whole-word matching on, these are the near misses."""
    if not text:
        return []
    flags = 0 if case_sensitive else re.IGNORECASE
    found = []
    for g in groups:
        for kw, _ in g.imap_patterns:
            if kw in found:
                continue
            if re.search(re.escape(kw), text, flags):
                found.append(kw)
                if len(found) >= limit:
                    return found
    return found


# ==========================================================================
# Document text extraction
# ==========================================================================
def extract_pdf(path):
    if not HAVE_PDF:
        raise RuntimeError("pypdf/PyPDF2 not installed")
    out = []
    reader = PdfReader(path)
    for page in reader.pages:
        try:
            out.append(page.extract_text() or "")
        except Exception:
            pass
    return "\n".join(out)


def extract_docx(path):
    if HAVE_DOCX:
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        # headers / footers
        for section in d.sections:
            for container in (section.header, section.footer):
                for p in container.paragraphs:
                    parts.append(p.text)
        return "\n".join(parts)
    if HAVE_DOCX2TXT:
        return docx2txt.process(path) or ""
    raise RuntimeError("python-docx / docx2txt not installed")


def extract_doc(path):
    # 1. antiword, if present
    if shutil.which("antiword"):
        try:
            res = subprocess.run(["antiword", path], capture_output=True, timeout=60)
            if res.returncode == 0 and res.stdout:
                return res.stdout.decode("utf-8", "ignore")
        except Exception:
            pass
    # 2. catdoc
    if shutil.which("catdoc"):
        try:
            res = subprocess.run(["catdoc", path], capture_output=True, timeout=60)
            if res.returncode == 0 and res.stdout:
                return res.stdout.decode("utf-8", "ignore")
        except Exception:
            pass
    # 3. olefile -> pull the WordDocument stream and strip
    if HAVE_OLEFILE:
        try:
            if olefile.isOleFile(path):
                ole = olefile.OleFileIO(path)
                if ole.exists("WordDocument"):
                    raw = ole.openstream("WordDocument").read()
                    ole.close()
                    return _strip_binary(raw)
                ole.close()
        except Exception:
            pass
    # 4. last-ditch raw scrape
    with open(path, "rb") as fh:
        return _strip_binary(fh.read())


def _strip_binary(raw: bytes) -> str:
    text = raw.decode("latin-1", "ignore")
    # keep printable runs, collapse the rest to whitespace
    text = re.sub(r"[^\x20-\x7E\r\n]+", " ", text)
    return re.sub(r"\s{2,}", " ", text)


def extract_text(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return extract_pdf(path)
    if ext == ".docx":
        return extract_docx(path)
    if ext == ".doc":
        return extract_doc(path)
    return ""


# ==========================================================================
# IMAP helpers
# ==========================================================================
def decode_hdr(value):
    if not value:
        return ""
    try:
        return str(make_header(decode_header(value)))
    except Exception:
        return value


def get_reply_to(msg):
    """Return (name, email) parsed from the message's Reply-To header,
    or (None, None) if there isn't one / it has no usable address."""
    if msg is None:
        return None, None
    raw = decode_hdr(msg.get("Reply-To", ""))
    if not raw:
        return None, None
    name, addr = email.utils.parseaddr(raw)
    addr = (addr or "").strip()
    name = (name or "").strip()
    if not addr:
        return None, None
    return (name or None), addr


def parse_date(value):
    if not value:
        return None
    value = value.strip()
    if HAVE_DATEUTIL:
        try:
            return dateparser.parse(value)
        except Exception:
            pass
    try:
        return email.utils.parsedate_to_datetime(value)
    except Exception:
        return None


class ImapClient:
    def __init__(self, server, port, username, password, use_ssl=True, mailbox="INBOX"):
        self.server = server
        self.port = int(port)
        self.username = username
        self.password = password
        self.use_ssl = use_ssl
        self.mailbox = mailbox or "INBOX"
        self.conn = None

    def connect(self):
        if self.use_ssl:
            self.conn = imaplib.IMAP4_SSL(self.server, self.port)
        else:
            self.conn = imaplib.IMAP4(self.server, self.port)
            try:
                self.conn.starttls()
            except Exception:
                pass
        self.conn.login(self.username, self.password)
        self.conn.select(self.mailbox, readonly=True)

    def close(self):
        if not self.conn:
            return
        try:
            self.conn.close()
        except Exception:
            pass
        try:
            self.conn.logout()
        except Exception:
            pass
        self.conn = None

    def list_mailboxes(self):
        typ, data = self.conn.list()
        boxes = []
        if typ == "OK":
            for raw in data:
                if not raw:
                    continue
                line = raw.decode("utf-8", "ignore")
                m = re.search(r'"([^"]+)"\s*$', line) or re.search(r"(\S+)\s*$", line)
                if m:
                    boxes.append(m.group(1))
        return boxes

    def search_message(self, subject, sender_email, date_obj, window_days=1):
        """Return a list of UIDs that plausibly match the CSV row."""
        criteria = []
        if date_obj:
            since = (date_obj - timedelta(days=window_days)).strftime("%d-%b-%Y")
            before = (date_obj + timedelta(days=window_days + 1)).strftime("%d-%b-%Y")
            criteria += ["SINCE", since, "BEFORE", before]
        if sender_email:
            criteria += ["FROM", f'"{sender_email}"']
        if subject:
            clean = subject.replace('"', " ").replace("\\", " ").strip()
            clean = re.sub(r"^(re|fwd|fw)\s*:\s*", "", clean, flags=re.I).strip()
            if clean:
                criteria += ["HEADER", "SUBJECT", f'"{clean[:120]}"']
        if not criteria:
            return []
        try:
            typ, data = self.conn.uid("SEARCH", None, *criteria)
        except Exception:
            return []
        if typ != "OK" or not data or not data[0]:
            return []
        return data[0].split()

    def fetch_message_header_date(self, uid):
        """Cheap fetch of just the Date: header, to disambiguate same-subject
        same-sender messages without pulling every header or the full body."""
        try:
            typ, data = self.conn.uid("FETCH", uid, "(BODY.PEEK[HEADER.FIELDS (DATE)])")
        except Exception:
            return None
        if typ != "OK" or not data:
            return None
        raw = b"".join(
            item[1] for item in data
            if isinstance(item, tuple) and len(item) >= 2
            and isinstance(item[1], bytes)
        )
        if not raw:
            return None
        text = raw.decode("utf-8", "ignore")
        m = re.search(r"Date:\s*(.+)", text, re.I)
        if not m:
            return None
        return parse_date(m.group(1).strip())

    def find_closest_message(self, uids, target_date, window_seconds=4, max_check=25):
        """Find the candidate whose Date header is closest to the CSV timestamp.
        Only checks up to max_check candidates to avoid excessive IMAP round
        trips when a sender/subject pair matches many messages on the same day."""
        if target_date is None:
            return None, None, None
        best_uid = best_date = best_delta = None
        for uid in uids[:max_check]:
            msg_date = self.fetch_message_header_date(uid)
            delta = datetime_delta_seconds(target_date, msg_date)
            if delta is None or delta > window_seconds:
                continue
            if best_delta is None or delta < best_delta:
                best_uid, best_date, best_delta = uid, msg_date, delta
        return best_uid, best_date, best_delta

    def fetch_message(self, uid):
        typ, data = self.conn.uid("FETCH", uid, "(RFC822)")
        if typ != "OK" or not data or not data[0]:
            return None
        return email.message_from_bytes(data[0][1])


class LocalEmlSource:
    """Drop-in replacement for ImapClient that reads a folder of .eml files
    instead of a mailbox - e.g. an extracted Zimbra tgz export, or any other
    export tool that leaves one standard .eml (RFC822) file per message.

    Implements the same methods Processor calls on ImapClient (connect,
    close, search_message, find_closest_message, fetch_message), so the rest
    of the pipeline - Reply-To, attachment extraction/exclusion, caching,
    keyword matching - needs no changes at all."""

    def __init__(self, folder, log_fn=None):
        self.folder = folder
        self._index = []   # list of dicts: path, from_header, subject_header, date
        self.log_fn = log_fn   # optional callable(str) for progress reporting

    def _log(self, msg):
        if self.log_fn:
            try:
                self.log_fn(msg)
            except Exception:
                pass

    def connect(self):
        if not self.folder or not os.path.isdir(self.folder):
            raise RuntimeError(f"Local export folder not found: {self.folder!r}")
        self._index = []
        scanned, skipped_other = 0, 0
        for root, _dirs, files in os.walk(self.folder):
            for fn in files:
                if os.path.splitext(fn)[1].lower() != ".eml":
                    skipped_other += 1
                    continue
                path = os.path.join(root, fn)
                try:
                    with open(path, "rb") as fh:
                        msg = email.parser.BytesParser(policy=email.policy.compat32)\
                            .parse(fh, headersonly=True)
                except Exception as exc:
                    self._log(f"    ! could not read header of {path}: {exc}")
                    continue
                self._index.append({
                    "path": path,
                    "from_header": decode_hdr(msg.get("From", "")),
                    "subject_header": decode_hdr(msg.get("Subject", "")),
                    "date": parse_date(msg.get("Date", "")),
                })
                scanned += 1
                if scanned % 1000 == 0:
                    self._log(f"    indexed {scanned} .eml files so far "
                             f"(currently in {os.path.relpath(root, self.folder)}) ...")
        if not self._index:
            raise RuntimeError(f"No .eml files found under {self.folder!r}")
        self._log(f"Indexed {len(self._index)} .eml files under {self.folder} "
                  f"({skipped_other} non-.eml file(s) skipped, e.g. .meta sidecars).")

    def close(self):
        self._index = []

    def search_message(self, subject, sender_email, date_obj, window_days=1):
        """Return a list of file paths (used as 'uids') that plausibly match
        the CSV row - same filtering logic as ImapClient.search_message, just
        applied to the in-memory index instead of an IMAP SEARCH command."""
        since = before = None
        if date_obj:
            since = date_obj - timedelta(days=window_days)
            before = date_obj + timedelta(days=window_days + 1)

        clean_subject = ""
        if subject:
            clean_subject = subject.replace('"', " ").strip()
            clean_subject = re.sub(r"^(re|fwd|fw)\s*:\s*", "", clean_subject, flags=re.I).strip()

        out = []
        for entry in self._index:
            if since is not None:
                if entry["date"] is None:
                    continue
                d = normalize_dt(entry["date"])
                if d is None or not (normalize_dt(since) <= d < normalize_dt(before)):
                    continue
            if sender_email and sender_email.lower() not in entry["from_header"].lower():
                continue
            if clean_subject and clean_subject.lower() not in entry["subject_header"].lower():
                continue
            out.append(entry["path"])
        return out

    def fetch_message_header_date(self, uid):
        """No extra I/O needed - the date was already read while indexing."""
        for entry in self._index:
            if entry["path"] == uid:
                return entry["date"]
        return None

    def find_closest_message(self, uids, target_date, window_seconds=4, max_check=25):
        """Same contract as ImapClient.find_closest_message, but free - the
        index already holds every candidate's parsed date, so there is no
        per-candidate fetch and no max_check truncation needed."""
        if target_date is None:
            return None, None, None
        best_uid = best_date = best_delta = None
        for uid in uids:
            msg_date = self.fetch_message_header_date(uid)
            delta = datetime_delta_seconds(target_date, msg_date)
            if delta is None or delta > window_seconds:
                continue
            if best_delta is None or delta < best_delta:
                best_uid, best_date, best_delta = uid, msg_date, delta
        return best_uid, best_date, best_delta

    def fetch_message(self, uid):
        try:
            with open(uid, "rb") as fh:
                return email.parser.BytesParser(policy=email.policy.compat32).parse(fh)
        except Exception:
            return None


def matches_any_pattern(filename, patterns):
    """True if filename matches any pattern in the list.
    Supports plain substrings ('signature.pdf', 'disclaimer') and glob
    wildcards ('*_signature.*', 'logo*.png'), all case-insensitive."""
    if not patterns:
        return False
    name = filename.lower()
    for pat in patterns:
        pat = (pat or "").strip().lower()
        if not pat:
            continue
        if fnmatch.fnmatch(name, pat):
            return True
        if pat in name:
            return True
    return False


# Kept as an alias so any external references to the old name still work.
is_excluded_attachment = matches_any_pattern


def save_attachments(msg, dest_dir, exclude_patterns=None, include_patterns=None):
    """Write allowed attachments to dest_dir, skipping ones that match
    exclude_patterns - UNLESS they also match include_patterns, which always
    wins (e.g. exclude 'notes' would normally drop
    'notes_reports_and_work.pdf', but including that exact name, or a
    wildcard like 'notes_reports*', forces it to download anyway).

    Returns (saved_paths, skipped_filenames, overridden_filenames)."""
    paths = []
    skipped = []
    overridden = []
    if msg is None:
        return paths, skipped, overridden
    for part in msg.walk():
        if part.get_content_maintype() == "multipart":
            continue
        disp = (part.get("Content-Disposition") or "").lower()
        filename = decode_hdr(part.get_filename())
        if not filename:
            continue
        if "attachment" not in disp and "inline" not in disp:
            continue
        ext = os.path.splitext(filename)[1].lower()
        if ext not in ALLOWED_EXTS:
            continue
        excluded = matches_any_pattern(filename, exclude_patterns)
        included = matches_any_pattern(filename, include_patterns)
        if excluded and included:
            overridden.append(filename)
        elif excluded:
            skipped.append(filename)
            continue
        safe = re.sub(r"[^A-Za-z0-9._-]+", "_", os.path.basename(filename))
        target = os.path.join(dest_dir, safe)
        n = 1
        while os.path.exists(target):
            stem, e = os.path.splitext(safe)
            target = os.path.join(dest_dir, f"{stem}_{n}{e}")
            n += 1
        try:
            payload = part.get_payload(decode=True)
            if not payload:
                continue
            with open(target, "wb") as fh:
                fh.write(payload)
            paths.append(target)
        except Exception:
            continue
    return paths, skipped, overridden


# ==========================================================================
# Processing engine (runs off the UI thread)
# ==========================================================================
class Processor(threading.Thread):
    def __init__(self, csv_path, out_dir, groups, options, imap_cfg, password, log_q,
                 exclude_patterns=None, include_patterns=None, local_cfg=None):
        super().__init__(daemon=True)
        self.csv_path = csv_path
        self.out_dir = out_dir
        self.groups = groups
        self.options = options
        self.imap_cfg = imap_cfg
        self.password = password
        self.log_q = log_q
        self.exclude_patterns = exclude_patterns or []
        self.include_patterns = include_patterns or []
        self.local_cfg = local_cfg or {}
        self.stop_flag = threading.Event()
        self.log_file_path = None
        self._log_fh = None

    def _explain_no_match(self, fname, text):
        """Log WHY an attachment produced no match (empty text / whole-word)."""
        if not (text or "").strip():
            self.log(f"    no text extracted from {fname} (scanned / image-only file?)", "warn")
            return
        if self.options.get("whole_word"):
            near = whole_word_near_misses(
                text, self.groups, self.options.get("case_sensitive", False))
            if near:
                self.log(f"    hint: {', '.join(repr(k) for k in near)} appear in {fname} only "
                         f"inside longer words - turn off 'Whole-word match' on the Groups "
                         f"tab to match them", "warn")

    def log(self, msg, kind="log"):
        self.log_q.put((kind, msg))
        if self._log_fh:
            ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            try:
                for line in str(msg).splitlines() or [""]:
                    self._log_fh.write(f"[{ts}] {line}\n")
                self._log_fh.flush()
            except Exception:
                pass

    # -- helpers ----------------------------------------------------------
    @staticmethod
    def _col(row, *names):
        for n in names:
            for key in row:
                if key and key.strip().lower() == n.lower():
                    return (row[key] or "").strip()
        return ""

    def run(self):
        logs_dir = os.path.join(self.out_dir, "logs")
        try:
            os.makedirs(logs_dir, exist_ok=True)
            stamp = datetime.now().strftime("%Y-%m-%d_%H%M%S")
            self.log_file_path = os.path.join(logs_dir, f"run_{stamp}.log")
            self._log_fh = open(self.log_file_path, "w", encoding="utf-8")
        except Exception as exc:
            self.log_file_path = None
            self._log_fh = None
            self.log(f"Could not open log file: {exc}", "warn")
        else:
            self.log(f"Logging to {self.log_file_path}")
            self.log_q.put(("log_file", self.log_file_path))

        try:
            self._run()
        except Exception:
            self.log(traceback.format_exc(), "error")
            self.log_q.put(("done", None))
        finally:
            if self._log_fh:
                try:
                    self._log_fh.close()
                except Exception:
                    pass
                self._log_fh = None

    def _run(self):
        for g in self.groups:
            g.compile(self.options["whole_word"], self.options["case_sensitive"])
        precedence = self.options.get("match_precedence", "keywords_first")
        self.log(f"Match precedence: {'group name first' if precedence == 'groupname_first' else 'keywords first'}")

        with open(self.csv_path, "r", encoding="utf-8-sig", newline="") as fh:
            rows = list(csv.DictReader(fh))
        total = len(rows)
        self.log(f"Loaded {total} rows from {os.path.basename(self.csv_path)}")
        self.log_q.put(("total", total))

        matched, unmatched = [], []
        needs_attachments = []

        # ---- Pass 1: subject + body ------------------------------------
        for idx, row in enumerate(rows, 1):
            if self.stop_flag.is_set():
                self.log("Stopped by user.", "warn")
                break
            name = self._col(row, "Sender", "Name", "From")
            addr = self._col(row, "Email", "Email Address", "From Email")
            date_s = self._col(row, "Date Received", "Date", "Received")
            subject = self._col(row, "Subject")
            body = self._col(row, "Body", "Message", "Content")

            group, term, stage = match_with_precedence(
                f"{subject}\n{body}", self.groups, precedence)
            if group:
                matched.append({
                    "Name": name,
                    "Email": addr,
                    "Descriptive Name": group.descriptive_name,
                    "Group Name": group.group_name,
                    "Date": date_s,
                    "Subject": subject,
                })
                self.log(f"[{idx}/{total}] {addr or '(no email)'} | {subject[:60]!r} "
                         f"-> {group.group_name} ({stage}: {term})")
            else:
                needs_attachments.append(
                    {"name": name, "email": addr, "date": date_s,
                     "subject": subject, "row": idx}
                )
            self.log_q.put(("progress", idx))

        self.log(f"Pass 1 complete: {len(matched)} matched, "
                 f"{len(needs_attachments)} need attachment search.")

        # ---- Pass 2: attachments - cache first, mail source only on a miss ----
        offline = bool(self.options.get("offline_mode"))
        use_cache = self.options.get("use_cache", True)
        cache = load_cache() if use_cache else {}
        cache_hits = 0
        n_attach = len(needs_attachments)

        source_kind = self.options.get("email_source", "imap")   # "imap" | "local"
        source_label = ("local export folder" if source_kind == "local" else "IMAP")

        if offline:
            searchable = bool(use_cache)
            if needs_attachments:
                if use_cache:
                    self.log(f"Offline mode: never touching the {source_label} - the remaining "
                             f"{n_attach} rows are matched against the local cache only.", "warn")
                else:
                    self.log("Offline mode with the cache disabled: nothing to search, "
                             f"{n_attach} rows go straight to unmatched.", "warn")
        else:
            searchable = bool(self.options.get("fetch_attachments"))

        def finish_unmatched(i, item, reason):
            unmatched.append({
                "Name": item["name"], "Email": item["email"],
                "Subject": item["subject"], "Date Received": item["date"],
                "Reason": reason,
            })
            self.log(f"[attach {i}/{n_attach}] {item['email'] or '(no email)'} | "
                     f"{item['subject'][:50]!r} -> UNMATCHED ({reason})")

        if needs_attachments and searchable and not self.stop_flag.is_set():
            client = None
            connect_error = None     # set once a connection attempt fails; not retried
            last_done = 0            # rows fully resolved so far (for the error handler)
            tmp_root = tempfile.mkdtemp(prefix="email_grouper_")
            try:
                window = int(self.options.get("date_window_days", 1))
                timestamp_window = float(self.options.get("timestamp_window_seconds", 4))
                for i, item in enumerate(needs_attachments, 1):
                    if self.stop_flag.is_set():
                        break
                    hit_group, hit_kw, reason = None, None, "no keyword or group name in attachment text"
                    out_name, out_email = item["name"], item["email"]
                    key = cache_key(item["email"], item["subject"], item["date"])
                    entry = cache.get(key) if use_cache else None

                    if entry is None:
                        # Cache miss: only now do we need the mail source (connect lazily).
                        if offline:
                            finish_unmatched(i, item, f"not in cache (offline mode - {source_label} skipped)")
                            last_done = i
                            continue
                        if client is None and connect_error is None:
                            try:
                                if source_kind == "local":
                                    folder = self.local_cfg.get("folder", "")
                                    self.log(f"Indexing local export folder: {folder} ...")
                                    client = LocalEmlSource(folder, log_fn=self.log)
                                else:
                                    self.log(f"Connecting to {self.imap_cfg['server']}:{self.imap_cfg['port']} ...")
                                    client = ImapClient(
                                        self.imap_cfg["server"], self.imap_cfg["port"],
                                        self.imap_cfg["username"], self.password,
                                        self.imap_cfg.get("use_ssl", True),
                                        self.imap_cfg.get("mailbox", "INBOX"),
                                    )
                                client.connect()
                                if source_kind == "imap":
                                    self.log("IMAP connected.")
                            except Exception as exc:
                                connect_error = exc
                                if client is not None:
                                    client.close()
                                client = None
                                self.log(f"{source_label} unavailable: {exc} - cached rows are "
                                         f"still matched, uncached rows go to unmatched.", "error")
                        if client is None:
                            finish_unmatched(i, item, f"{source_label} unavailable: {connect_error}")
                            last_done = i
                            continue

                    if entry is not None:
                        # ---- served entirely from cache, no server hit ----
                        cache_hits += 1
                        atts = entry.get("attachments", [])
                        if entry.get("reply_email"):
                            out_name = entry.get("reply_name") or out_name
                            out_email = entry["reply_email"]
                        self.log(f"[attach {i}/{len(needs_attachments)}] "
                                 f"{out_email} - {len(atts)} attachment(s) "
                                 f"(cached, no server call)")
                        if entry.get("not_found"):
                            reason = entry.get("reason") or "message not found on IMAP server (cached)"
                        elif not atts:
                            reason = "no pdf/docx/doc attachments (cached)"
                        for j, att in enumerate(atts, 1):
                            fname = att.get("filename", "?")
                            self.log(f"    parsing cached attachment {j}/{len(atts)}: {fname}")
                            g, kw, stage = match_with_precedence(
                                att.get("text", ""), self.groups, precedence, scope="imap")
                            if g:
                                hit_group, hit_kw = g, kw
                                self.log(f"    + {fname} -> {g.group_name} ({stage}: {kw})")
                                break
                            self._explain_no_match(fname, att.get("text", ""))
                    else:
                        # ---- go to the server ----
                        date_obj = parse_date(item["date"])
                        uids = client.search_message(item["subject"], item["email"], date_obj, window)
                        cache_atts = []
                        if not uids:
                            reason = f"message not found on {source_label}"
                            self.log(f"[attach {i}/{len(needs_attachments)}] "
                                     f"{item['email']} - message not found on {source_label}")
                            if use_cache:
                                cache[key] = {"not_found": True, "attachments": [], "reason": reason}
                        elif date_obj is None:
                            reason = "CSV date/time could not be parsed; exact timestamp match skipped"
                            self.log(f"[attach {i}/{len(needs_attachments)}] "
                                     f"{item['email']} - cannot parse CSV timestamp; refusing to guess a message")
                            if use_cache:
                                cache[key] = {"not_found": True, "attachments": [], "reason": reason}
                        else:
                            if len(uids) > 25:
                                self.log(f"    {len(uids)} same subject/sender candidates found; "
                                         f"checking timestamps on first 25 only", "warn")
                            uid, matched_date, delta = client.find_closest_message(
                                uids, date_obj, timestamp_window, max_check=25)
                            if uid is None:
                                reason = f"no IMAP message within ±{timestamp_window:g} seconds of CSV timestamp"
                                self.log(f"[attach {i}/{len(needs_attachments)}] "
                                         f"{item['email']} - no timestamp match within ±{timestamp_window:g}s")
                                if use_cache:
                                    cache[key] = {"not_found": True, "attachments": [], "reason": reason}
                            else:
                                msg = client.fetch_message(uid)
                                reply_name, reply_email = get_reply_to(msg)
                                if reply_email:
                                    out_name = reply_name or out_name
                                    out_email = reply_email
                                    self.log(f"    Reply-To found: using "
                                             f"{out_name or '(no name)'} <{out_email}> "
                                             f"instead of CSV sender")
                                work = tempfile.mkdtemp(dir=tmp_root)
                                files, skipped, overridden = save_attachments(
                                    msg, work,
                                    exclude_patterns=self.exclude_patterns,
                                    include_patterns=self.include_patterns)
                                self.log(f"[attach {i}/{len(needs_attachments)}] "
                                         f"{out_email} - matched IMAP message {matched_date} "
                                         f"({delta:g}s from CSV), {len(files)} attachment(s) found"
                                         + (f", {len(skipped)} excluded" if skipped else ""))
                                if overridden:
                                    self.log(f"    included despite exclude filter: "
                                             f"{', '.join(overridden)}")
                                if skipped:
                                    self.log(f"    excluded: {', '.join(skipped)}")
                                if not files:
                                    reason = "no pdf/docx/doc attachments" if not skipped \
                                        else "all attachments excluded by filter"
                                for j, path in enumerate(files, 1):
                                    fname = os.path.basename(path)
                                    self.log(f"    parsing attachment {j}/{len(files)}: {fname}")
                                    try:
                                        text = extract_text(path)
                                    except Exception as exc:
                                        self.log(f"    ! {fname}: {exc}", "warn")
                                        text = ""
                                    cache_atts.append({"filename": fname, "text": text})
                                    if not hit_group:
                                        g, kw, stage = match_with_precedence(
                                            text, self.groups, precedence, scope="imap")
                                        if g:
                                            hit_group, hit_kw = g, kw
                                            self.log(f"    + {fname} -> {g.group_name} ({stage}: {kw})")
                                        else:
                                            self._explain_no_match(fname, text)
                                if use_cache:
                                    cache[key] = {
                                        "attachments": cache_atts,
                                        "message_uid": uid.decode(errors="replace") if isinstance(uid, bytes) else str(uid),
                                        "message_date": matched_date.isoformat() if matched_date else "",
                                        "timestamp_delta_seconds": delta,
                                        "reply_name": reply_name or "",
                                        "reply_email": reply_email or "",
                                    }
                        if use_cache:
                            save_cache(cache)

                    if hit_group:
                        matched.append({
                            "Name": out_name,
                            "Email": out_email,
                            "Descriptive Name": hit_group.descriptive_name,
                            "Group Name": hit_group.group_name,
                            "Date": item["date"],
                            "Subject": item["subject"],
                        })
                    else:
                        unmatched.append({
                            "Name": out_name,
                            "Email": out_email,
                            "Subject": item["subject"],
                            "Date Received": item["date"],
                            "Reason": reason,
                        })
                    self.log(f"[attach {i}/{len(needs_attachments)}] {out_email or '(no email)'} | "
                             f"{item['subject'][:50]!r} "
                             f"-> {hit_group.group_name if hit_group else 'UNMATCHED'}")
                    last_done = i
                if use_cache:
                    self.log(f"Cache: {cache_hits}/{len(needs_attachments)} lookups served "
                             f"from cache ({CACHE_PATH}).")
            except Exception as exc:
                self.log(f"{source_label} error: {exc}", "error")
                # Only rows not already resolved above (no double-counting).
                for item in needs_attachments[last_done:]:
                    unmatched.append({
                        "Name": item["name"], "Email": item["email"],
                        "Subject": item["subject"], "Date Received": item["date"],
                        "Reason": f"{source_label} error during attachment search: {exc}",
                    })
            finally:
                if client:
                    client.close()
                shutil.rmtree(tmp_root, ignore_errors=True)
        else:
            if self.stop_flag.is_set():
                reason = "run stopped before attachment search"
            elif offline:
                reason = "no keyword in subject/body (offline mode, cache disabled)"
            else:
                reason = "no keyword in subject/body (attachment search disabled)"
            for item in needs_attachments:
                unmatched.append({
                    "Name": item["name"], "Email": item["email"],
                    "Subject": item["subject"], "Date Received": item["date"],
                    "Reason": reason,
                })

        # ---- Export ----------------------------------------------------
        matched_path = os.path.join(self.out_dir, "matched.csv")
        unmatched_path = os.path.join(self.out_dir, "unmatched.csv")
        fieldnames = ["Name", "Email", "Descriptive Name", "Group Name", "Date", "Subject"]

        with open(matched_path, "w", encoding="utf-8", newline="") as fh:
            w = csv.DictWriter(fh, fieldnames=fieldnames)
            w.writeheader()
            w.writerows(matched)

        with open(unmatched_path, "w", encoding="utf-8", newline="") as fh:
            w = csv.DictWriter(
                fh, fieldnames=["Name", "Email", "Subject", "Date Received", "Reason"])
            w.writeheader()
            w.writerows(unmatched)

        self.log(f"Wrote {len(matched)} rows -> {matched_path}", "ok")
        self.log(f"Wrote {len(unmatched)} rows -> {unmatched_path}", "ok")

        if self.options.get("dedupe_by_email"):
            seen = set()
            matched_uniq = []
            for row in matched:
                dkey = (row["Group Name"], (row["Email"] or "").strip().lower())
                if dkey in seen:
                    continue
                seen.add(dkey)
                matched_uniq.append(row)
            uniq_path = os.path.join(self.out_dir, "matched_uniq.csv")
            with open(uniq_path, "w", encoding="utf-8", newline="") as fh:
                w = csv.DictWriter(fh, fieldnames=fieldnames)
                w.writeheader()
                w.writerows(matched_uniq)
            removed = len(matched) - len(matched_uniq)
            self.log(f"Wrote {len(matched_uniq)} unique-per-group rows "
                     f"({removed} duplicate email(s) filtered out) -> {uniq_path}", "ok")

            # Unmatched rows have no group, so dedupe is by email alone.
            seen_u = set()
            unmatched_uniq = []
            for row in unmatched:
                dkey = (row["Email"] or "").strip().lower()
                if dkey in seen_u:
                    continue
                seen_u.add(dkey)
                unmatched_uniq.append(row)
            unmatched_uniq_path = os.path.join(self.out_dir, "unmatched_uniq.csv")
            with open(unmatched_uniq_path, "w", encoding="utf-8", newline="") as fh:
                w = csv.DictWriter(
                    fh, fieldnames=["Name", "Email", "Subject", "Date Received", "Reason"])
                w.writeheader()
                w.writerows(unmatched_uniq)
            removed_u = len(unmatched) - len(unmatched_uniq)
            self.log(f"Wrote {len(unmatched_uniq)} unique unmatched rows "
                     f"({removed_u} duplicate email(s) filtered out) -> {unmatched_uniq_path}", "ok")

        self.log_q.put(("done", (len(matched), len(unmatched))))


# ==========================================================================
# GUI
# ==========================================================================
class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Email Grouper")
        self.geometry("980x680")
        self.minsize(860, 600)

        self.cfg = load_config()
        self.log_q = queue.Queue()
        self.processor = None

        self.csv_path = tk.StringVar(value=self.cfg.get("recent", {}).get("last_csv", ""))
        self.out_dir = tk.StringVar(
            value=self.cfg.get("recent", {}).get("last_out_dir", "") or APP_DIR)

        nb = ttk.Notebook(self)
        nb.pack(fill="both", expand=True, padx=8, pady=8)
        self.tab_groups = ttk.Frame(nb)
        self.tab_imap = ttk.Frame(nb)
        self.tab_excl = ttk.Frame(nb)
        self.tab_run = ttk.Frame(nb)
        nb.add(self.tab_groups, text="Groups & Keywords")
        nb.add(self.tab_imap, text="IMAP Login")
        nb.add(self.tab_excl, text="Attachment Filters")
        nb.add(self.tab_run, text="Run")

        self._build_groups_tab()
        self._build_imap_tab()
        self._build_excl_tab()
        self._build_run_tab()

        self._refresh_group_tree()
        self._toggle_offline()
        self._toggle_source()
        self.after(150, self._drain_log)
        self.protocol("WM_DELETE_WINDOW", self._on_close)

    # ---------------- Groups tab ----------------------------------------
    def _build_groups_tab(self):
        f = self.tab_groups
        left = ttk.Frame(f)
        left.pack(side="left", fill="both", expand=True, padx=(10, 5), pady=10)
        right = ttk.LabelFrame(f, text="Group editor")
        right.pack(side="right", fill="y", padx=(5, 10), pady=10)

        cols = ("group", "descriptive", "imap", "body")
        self.tree = ttk.Treeview(left, columns=cols, show="headings", selectmode="browse")
        self.tree.heading("group", text="Group name")
        self.tree.heading("descriptive", text="Descriptive name (line 1)")
        self.tree.heading("imap", text="Line 2 kw (attach + subj/body)")
        self.tree.heading("body", text="Lines 3+ kw (subj/body)")
        self.tree.column("group", width=130)
        self.tree.column("descriptive", width=110)
        self.tree.column("imap", width=200)
        self.tree.column("body", width=200)
        vs = ttk.Scrollbar(left, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=vs.set)
        self.tree.pack(side="left", fill="both", expand=True)
        vs.pack(side="right", fill="y")
        self.tree.bind("<<TreeviewSelect>>", self._on_group_select)

        ttk.Label(right, text="Group name").pack(anchor="w", padx=10, pady=(10, 0))
        self.e_group = ttk.Entry(right, width=34)
        self.e_group.pack(padx=10, pady=2)
        ttk.Label(right, text="e.g. G1/1/1/2026/01",
                  foreground="#777").pack(anchor="w", padx=10)

        ttk.Label(right, text="Keywords").pack(anchor="w", padx=10, pady=(12, 0))
        ttk.Label(right, foreground="#777", justify="left",
                  text=("Line 1: descriptive name (shown in the CSVs). Also\n"
                        "matched against attachments AND subject/body.\n"
                        "Line 2: comma-separated keywords, matched against\n"
                        "attachments AND subject/body.\n"
                        "   e.g.  jan, january, jaanuary, njan\n"
                        "Lines 3+: one keyword per line, subject/body only."))\
            .pack(anchor="w", padx=10)
        self.t_keywords = tk.Text(right, width=34, height=12, wrap="word")
        self.t_keywords.pack(padx=10, pady=4)

        btns = ttk.Frame(right)
        btns.pack(fill="x", padx=10, pady=8)
        ttk.Button(btns, text="Add / Update", command=self._add_group).pack(fill="x", pady=2)
        ttk.Button(btns, text="Delete selected", command=self._delete_group).pack(fill="x", pady=2)
        ttk.Button(btns, text="Clear form", command=self._clear_group_form).pack(fill="x", pady=2)
        ttk.Separator(right).pack(fill="x", padx=10, pady=6)

        opt = ttk.LabelFrame(right, text="Matching options")
        opt.pack(fill="x", padx=10, pady=4)
        self.v_whole = tk.BooleanVar(value=self.cfg["options"]["whole_word"])
        self.v_case = tk.BooleanVar(value=self.cfg["options"]["case_sensitive"])
        ttk.Checkbutton(opt, text="Whole-word match", variable=self.v_whole).pack(anchor="w", padx=8, pady=2)
        ttk.Checkbutton(opt, text="Case sensitive", variable=self.v_case).pack(anchor="w", padx=8, pady=2)

        ttk.Separator(opt, orient="horizontal").pack(fill="x", padx=8, pady=4)
        ttk.Label(opt, text="Match precedence").pack(anchor="w", padx=8)
        self.v_precedence = tk.StringVar(
            value=self.cfg["options"].get("match_precedence", "keywords_first"))
        ttk.Radiobutton(opt, text="Keywords first, then group name",
                        variable=self.v_precedence, value="keywords_first")\
            .pack(anchor="w", padx=8, pady=1)
        ttk.Radiobutton(opt, text="Group name first, then keywords",
                        variable=self.v_precedence, value="groupname_first")\
            .pack(anchor="w", padx=8, pady=1)

        ttk.Button(right, text="Save config.json",
                   command=self._save_all).pack(fill="x", padx=10, pady=(8, 12))

    def _refresh_group_tree(self):
        self.tree.delete(*self.tree.get_children())
        for i, g in enumerate(self.cfg["groups"]):
            descriptive, csv_kws, extra_kws = split_keyword_lines(g.get("keywords", []))
            self.tree.insert("", "end", iid=str(i), values=(
                g.get("group_name", ""),
                descriptive,
                ", ".join(csv_kws),
                " | ".join(extra_kws),
            ))

    def _on_group_select(self, _event=None):
        sel = self.tree.selection()
        if not sel:
            return
        g = self.cfg["groups"][int(sel[0])]
        self.e_group.delete(0, "end")
        self.e_group.insert(0, g.get("group_name", ""))
        self.t_keywords.delete("1.0", "end")
        self.t_keywords.insert("1.0", "\n".join(g.get("keywords", [])))

    def _clear_group_form(self):
        self.e_group.delete(0, "end")
        self.t_keywords.delete("1.0", "end")
        self.tree.selection_remove(self.tree.selection())

    def _add_group(self):
        name = self.e_group.get().strip()
        kws = [l.strip() for l in self.t_keywords.get("1.0", "end").splitlines() if l.strip()]
        if not name:
            messagebox.showwarning("Missing", "Group name is required.")
            return
        if not kws:
            messagebox.showwarning("Missing", "Enter the keywords: line 1 = descriptive "
                                              "name, line 2 = comma-separated keywords "
                                              "to match.")
            return
        entry = {"group_name": name, "keywords": kws}
        for i, g in enumerate(self.cfg["groups"]):
            if g.get("group_name") == name:
                self.cfg["groups"][i] = entry
                break
        else:
            self.cfg["groups"].append(entry)
        self._refresh_group_tree()
        self._clear_group_form()

    def _delete_group(self):
        sel = self.tree.selection()
        if not sel:
            return
        idx = int(sel[0])
        name = self.cfg["groups"][idx].get("group_name", "")
        if messagebox.askyesno("Delete", f"Delete group {name!r}?"):
            del self.cfg["groups"][idx]
            self._refresh_group_tree()
            self._clear_group_form()

    # ---------------- Attachment Filters tab -----------------------------
    def _build_excl_tab(self):
        f = self.tab_excl
        intro = ("Attachments matching any EXCLUDE pattern below are skipped entirely - "
                 "never downloaded, never opened, never searched for keywords.\n"
                 "The INCLUDE list overrides excludes: if a filename matches both lists, "
                 "it is downloaded anyway. Example: exclude 'notes' would normally drop "
                 "'notes_reports_and_work.pdf', but adding that name (or a pattern like "
                 "'notes_reports*') to the include list forces it through.\n"
                 "Matching is case-insensitive. Use a plain filename or fragment "
                 "(e.g. 'disclaimer.pdf', 'signature') for a substring match, or a "
                 "wildcard pattern (e.g. '*_logo.*', 'terms*.docx') for glob matching.")
        ttk.Label(f, text=intro, wraplength=820, justify="left")\
            .pack(anchor="w", padx=14, pady=(14, 8))

        # ---- Exclude list ----
        excl_frame = ttk.LabelFrame(f, text="Exclude list")
        excl_frame.pack(fill="both", expand=True, padx=14, pady=(0, 6))
        self.excl_listbox = self._build_pattern_panel(
            excl_frame, self.cfg.get("excluded_attachments", []))

        # ---- Include list (overrides excludes) ----
        incl_frame = ttk.LabelFrame(f, text="Include list (overrides excludes)")
        incl_frame.pack(fill="both", expand=True, padx=14, pady=(6, 8))
        self.incl_listbox = self._build_pattern_panel(
            incl_frame, self.cfg.get("included_attachments", []))

        ttk.Button(f, text="Save config.json", command=self._save_all)\
            .pack(anchor="e", padx=14, pady=(0, 12))

    def _build_pattern_panel(self, parent, initial_patterns):
        """Builds a listbox + entry + add/remove/clear controls inside `parent`.
        Returns the Listbox widget."""
        body = ttk.Frame(parent)
        body.pack(fill="both", expand=True, padx=8, pady=8)

        left = ttk.Frame(body)
        left.pack(side="left", fill="both", expand=True)
        listbox = tk.Listbox(left, height=6, selectmode="extended")
        vs = ttk.Scrollbar(left, orient="vertical", command=listbox.yview)
        listbox.configure(yscrollcommand=vs.set)
        listbox.pack(side="left", fill="both", expand=True)
        vs.pack(side="right", fill="y")
        for pat in initial_patterns:
            listbox.insert("end", pat)

        right = ttk.Frame(body)
        right.pack(side="left", fill="y", padx=(12, 0))
        entry = ttk.Entry(right, width=30)
        entry.pack(pady=(0, 4))
        entry.bind("<Return>", lambda _e: self._add_pattern(listbox, entry))
        ttk.Button(right, text="Add", command=lambda: self._add_pattern(listbox, entry))\
            .pack(fill="x", pady=2)
        ttk.Button(right, text="Remove selected",
                   command=lambda: self._remove_pattern(listbox)).pack(fill="x", pady=2)
        ttk.Button(right, text="Clear all",
                   command=lambda: self._clear_patterns(listbox)).pack(fill="x", pady=2)
        return listbox

    @staticmethod
    def _add_pattern(listbox, entry):
        pat = entry.get().strip()
        if not pat:
            return
        existing = [p.lower() for p in listbox.get(0, "end")]
        if pat.lower() not in existing:
            listbox.insert("end", pat)
        entry.delete(0, "end")

    @staticmethod
    def _remove_pattern(listbox):
        for idx in reversed(listbox.curselection()):
            listbox.delete(idx)

    @staticmethod
    def _clear_patterns(listbox):
        if listbox.size() and messagebox.askyesno("Clear all", "Remove all patterns in this list?"):
            listbox.delete(0, "end")

    # ---------------- IMAP tab ------------------------------------------
    def _build_imap_tab(self):
        f = ttk.LabelFrame(self.tab_imap, text="IMAP server")
        f.pack(fill="x", padx=14, pady=14)

        im = self.cfg["imap"]
        self.v_server = tk.StringVar(value=im.get("server", ""))
        self.v_port = tk.StringVar(value=str(im.get("port", 993)))
        self.v_user = tk.StringVar(value=im.get("username", ""))
        self.v_ssl = tk.BooleanVar(value=im.get("use_ssl", True))
        self.v_mailbox = tk.StringVar(value=im.get("mailbox", "INBOX"))
        self.v_pass = tk.StringVar()

        rows = [
            ("Server", self.v_server, False),
            ("Port", self.v_port, False),
            ("Username", self.v_user, False),
            ("Password (not saved)", self.v_pass, True),
            ("Mailbox", self.v_mailbox, False),
        ]
        for r, (label, var, secret) in enumerate(rows):
            ttk.Label(f, text=label).grid(row=r, column=0, sticky="w", padx=10, pady=5)
            e = ttk.Entry(f, textvariable=var, width=44, show="*" if secret else "")
            e.grid(row=r, column=1, sticky="w", padx=10, pady=5)

        ttk.Checkbutton(f, text="Use SSL (port 993)", variable=self.v_ssl)\
            .grid(row=len(rows), column=1, sticky="w", padx=10, pady=4)

        bar = ttk.Frame(self.tab_imap)
        bar.pack(fill="x", padx=14)
        ttk.Button(bar, text="Test connection", command=self._test_imap).pack(side="left", padx=4)
        ttk.Button(bar, text="Save (no password)", command=self._save_all).pack(side="left", padx=4)

        note = ("The password is held in memory for this session only and is never\n"
                "written to config.json.")
        ttk.Label(self.tab_imap, text=note, foreground="#777")\
            .pack(anchor="w", padx=18, pady=10)

        dep = ttk.LabelFrame(self.tab_imap, text="Attachment parsers detected")
        dep.pack(fill="x", padx=14, pady=8)
        lines = [
            f"PDF  (pypdf/PyPDF2)      : {'yes' if HAVE_PDF else 'no  - pip install pypdf'}",
            f"DOCX (python-docx)       : {'yes' if HAVE_DOCX else 'no  - pip install python-docx'}",
            f"DOCX fallback (docx2txt) : {'yes' if HAVE_DOCX2TXT else 'no  - pip install docx2txt'}",
            f"DOC  (antiword)          : {'yes' if shutil.which('antiword') else 'no  - apt install antiword'}",
            f"DOC  fallback (olefile)  : {'yes' if HAVE_OLEFILE else 'no  - pip install olefile'}",
            f"Date parsing (dateutil)  : {'yes' if HAVE_DATEUTIL else 'no  - pip install python-dateutil'}",
        ]
        for ln in lines:
            ttk.Label(dep, text=ln, font=("TkFixedFont", 9)).pack(anchor="w", padx=10, pady=1)

    def _imap_cfg(self):
        try:
            port = int(self.v_port.get())
        except ValueError:
            port = 993
        return {
            "server": self.v_server.get().strip(),
            "port": port,
            "username": self.v_user.get().strip(),
            "use_ssl": self.v_ssl.get(),
            "mailbox": self.v_mailbox.get().strip() or "INBOX",
        }

    def _local_cfg(self):
        return {"folder": self.v_local_folder.get().strip()}

    def _test_imap(self):
        cfg = self._imap_cfg()
        if not cfg["server"] or not cfg["username"]:
            messagebox.showwarning("Missing", "Server and username are required.")
            return
        if not self.v_pass.get():
            messagebox.showwarning("Missing", "Enter the password for this session.")
            return

        def work():
            c = ImapClient(cfg["server"], cfg["port"], cfg["username"],
                           self.v_pass.get(), cfg["use_ssl"], cfg["mailbox"])
            try:
                c.connect()
                boxes = c.list_mailboxes()
                self.log_q.put(("ok", f"Connected. {len(boxes)} mailboxes: "
                                      f"{', '.join(boxes[:12])}"))
            except Exception as exc:
                self.log_q.put(("error", f"Connection failed: {exc}"))
            finally:
                c.close()

        self._log("Testing IMAP connection ...")
        threading.Thread(target=work, daemon=True).start()

    # ---------------- Run tab -------------------------------------------
    def _build_run_tab(self):
        f = self.tab_run
        top = ttk.LabelFrame(f, text="Input / output")
        top.pack(fill="x", padx=14, pady=(14, 6))

        ttk.Label(top, text="Emails CSV").grid(row=0, column=0, sticky="w", padx=10, pady=6)
        ttk.Entry(top, textvariable=self.csv_path, width=68)\
            .grid(row=0, column=1, padx=6, pady=6)
        ttk.Button(top, text="Browse...", command=self._pick_csv)\
            .grid(row=0, column=2, padx=8)

        ttk.Label(top, text="Output folder").grid(row=1, column=0, sticky="w", padx=10, pady=6)
        ttk.Entry(top, textvariable=self.out_dir, width=68)\
            .grid(row=1, column=1, padx=6, pady=6)
        ttk.Button(top, text="Browse...", command=self._pick_out)\
            .grid(row=1, column=2, padx=8)

        self.v_fetch = tk.BooleanVar(value=self.cfg["options"]["fetch_attachments"])
        self.v_offline = tk.BooleanVar(value=self.cfg["options"].get("offline_mode", False))
        self.v_window = tk.StringVar(value=str(self.cfg["options"]["date_window_days"]))
        self.v_timestamp_window = tk.StringVar(
            value=str(self.cfg["options"].get("timestamp_window_seconds", 4))
        )

        opts = ttk.Frame(top)
        opts.grid(row=2, column=1, sticky="w", padx=6, pady=4)
        self.cb_fetch = ttk.Checkbutton(
            opts, text="Search attachments over IMAP when subject/body miss",
            variable=self.v_fetch)
        self.cb_fetch.pack(side="left")
        ttk.Label(opts, text="   search date window ±").pack(side="left")
        self.e_window = ttk.Entry(opts, textvariable=self.v_window, width=4)
        self.e_window.pack(side="left")
        ttk.Label(opts, text="days").pack(side="left")
        ttk.Label(opts, text="   exact timestamp ±").pack(side="left")
        self.e_timestamp_window = ttk.Entry(
            opts, textvariable=self.v_timestamp_window, width=5)
        self.e_timestamp_window.pack(side="left")
        ttk.Label(opts, text="seconds").pack(side="left")

        source_frame = ttk.LabelFrame(top, text="Mail source (used when attachment search is needed)")
        source_frame.grid(row=3, column=1, sticky="we", padx=6, pady=(0, 6))
        self.v_source = tk.StringVar(value=self.cfg["options"].get("email_source", "imap"))
        src_row = ttk.Frame(source_frame)
        src_row.pack(fill="x", padx=8, pady=(6, 2))
        ttk.Radiobutton(src_row, text="IMAP server (IMAP Login tab)", variable=self.v_source,
                        value="imap", command=self._toggle_source).pack(side="left")
        ttk.Radiobutton(src_row, text="Local export folder (.eml files, "
                                      "e.g. extracted Zimbra tgz)",
                        variable=self.v_source, value="local",
                        command=self._toggle_source).pack(side="left", padx=(16, 0))
        local_row = ttk.Frame(source_frame)
        local_row.pack(fill="x", padx=8, pady=(2, 8))
        self.v_local_folder = tk.StringVar(value=self.cfg.get("local", {}).get("folder", ""))
        ttk.Label(local_row, text="Folder:").pack(side="left")
        self.e_local_folder = ttk.Entry(local_row, textvariable=self.v_local_folder, width=52)
        self.e_local_folder.pack(side="left", padx=6)
        self.btn_local_browse = ttk.Button(local_row, text="Browse...", command=self._pick_local_folder)
        self.btn_local_browse.pack(side="left")
        ttk.Label(source_frame,
                  text="Scanned recursively for *.eml files (subfolders like Inbox!1, "
                       "Inbox!2 ... are picked up automatically).",
                  foreground="#777", wraplength=560, justify="left")\
            .pack(anchor="w", padx=8, pady=(0, 6))

        over = ttk.Frame(top)
        over.grid(row=4, column=1, sticky="w", padx=6, pady=(0, 4))
        ttk.Checkbutton(
            over,
            text="OVERRIDE: offline mode - never touch the mail source; match from the "
                 "local cache only, anything not cached is UNMATCHED",
            variable=self.v_offline, command=self._toggle_offline).pack(side="left")

        cache_row = ttk.Frame(top)
        cache_row.grid(row=5, column=1, sticky="w", padx=6, pady=(0, 8))
        self.v_use_cache = tk.BooleanVar(value=self.cfg["options"].get("use_cache", True))
        ttk.Checkbutton(cache_row, text="Use email/attachment cache "
                                        "(skip server for already-looked-up messages)",
                        variable=self.v_use_cache).pack(side="left")
        ttk.Button(cache_row, text="Clear cache", command=self._clear_cache).pack(side="left", padx=8)

        dedupe_row = ttk.Frame(top)
        dedupe_row.grid(row=6, column=1, sticky="w", padx=6, pady=(0, 8))
        self.v_dedupe = tk.BooleanVar(value=self.cfg["options"].get("dedupe_by_email", False))
        ttk.Checkbutton(
            dedupe_row,
            text="Also export matched_uniq.csv + unmatched_uniq.csv (unique by email)",
            variable=self.v_dedupe).pack(side="left")

        self.l_cache_info = ttk.Label(top, foreground="#777")
        self.l_cache_info.grid(row=7, column=1, sticky="w", padx=6, pady=(0, 6))
        self._refresh_cache_info()

        bar = ttk.Frame(f)
        bar.pack(fill="x", padx=14, pady=6)
        self.btn_run = ttk.Button(bar, text="Run", command=self._run)
        self.btn_run.pack(side="left", padx=4)
        self.btn_stop = ttk.Button(bar, text="Stop", command=self._stop, state="disabled")
        self.btn_stop.pack(side="left", padx=4)
        self.progress = ttk.Progressbar(bar, mode="determinate", length=460)
        self.progress.pack(side="left", padx=14)

        logf = ttk.LabelFrame(f, text="Log")
        logf.pack(fill="both", expand=True, padx=14, pady=(6, 14))
        self.log = tk.Text(logf, height=18, wrap="word", state="disabled")
        ls = ttk.Scrollbar(logf, orient="vertical", command=self.log.yview)
        self.log.configure(yscrollcommand=ls.set)
        self.log.pack(side="left", fill="both", expand=True)
        ls.pack(side="right", fill="y")
        self.log.tag_configure("error", foreground="#c0392b")
        self.log.tag_configure("warn", foreground="#b9770e")
        self.log.tag_configure("ok", foreground="#1e8449")

    def _toggle_offline(self):
        state = "disabled" if self.v_offline.get() else "normal"
        self.cb_fetch.configure(state=state)
        self.e_window.configure(state=state)
        self.e_timestamp_window.configure(state=state)

    def _toggle_source(self):
        state = "normal" if self.v_source.get() == "local" else "disabled"
        self.e_local_folder.configure(state=state)
        self.btn_local_browse.configure(state=state)

    def _pick_local_folder(self):
        p = filedialog.askdirectory(
            title="Select extracted mail export folder",
            initialdir=self.v_local_folder.get() or None)
        if p:
            self.v_local_folder.set(p)
            self.cfg.setdefault("local", {})["folder"] = p
            self._save_recent()

    def _refresh_cache_info(self):
        if os.path.exists(CACHE_PATH):
            try:
                n = len(load_cache())
            except Exception:
                n = "?"
            self.l_cache_info.configure(
                text=f"Cache file: {CACHE_PATH}  ({n} messages cached)")
        else:
            self.l_cache_info.configure(text=f"Cache file: {CACHE_PATH}  (empty)")

    def _clear_cache(self):
        if os.path.exists(CACHE_PATH):
            if not messagebox.askyesno("Clear cache", "Delete the cached IMAP lookups?"):
                return
            os.remove(CACHE_PATH)
        self._refresh_cache_info()
        self._log("Attachment cache cleared.", "ok")

    def _pick_csv(self):
        p = filedialog.askopenfilename(
            title="Select emails CSV",
            initialdir=os.path.dirname(self.csv_path.get()) if self.csv_path.get() else None,
            filetypes=[("CSV files", "*.csv"), ("All files", "*.*")])
        if p:
            self.csv_path.set(p)
            if not self.out_dir.get():
                self.out_dir.set(os.path.dirname(p))
            self.cfg.setdefault("recent", {})["last_csv"] = p
            self._save_recent()

    def _pick_out(self):
        p = filedialog.askdirectory(
            title="Select output folder",
            initialdir=self.out_dir.get() or None)
        if p:
            self.out_dir.set(p)
            self.cfg.setdefault("recent", {})["last_out_dir"] = p
            self._save_recent()

    def _save_recent(self):
        try:
            save_config(self.cfg)
        except Exception:
            pass

    # ---------------- shared --------------------------------------------
    def _collect_options(self):
        try:
            window = max(0, int(self.v_window.get()))
        except ValueError:
            window = 1
        try:
            timestamp_window = max(0.0, float(self.v_timestamp_window.get()))
        except ValueError:
            timestamp_window = 4.0
        return {
            "whole_word": self.v_whole.get(),
            "case_sensitive": self.v_case.get(),
            "match_precedence": self.v_precedence.get(),
            "fetch_attachments": self.v_fetch.get(),
            "offline_mode": self.v_offline.get(),
            "use_cache": self.v_use_cache.get(),
            "dedupe_by_email": self.v_dedupe.get(),
            "date_window_days": window,
            "timestamp_window_seconds": timestamp_window,
            "email_source": self.v_source.get(),
        }

    def _save_all(self):
        self.cfg["imap"] = self._imap_cfg()
        self.cfg["local"] = self._local_cfg()
        self.cfg["options"] = self._collect_options()
        self.cfg["excluded_attachments"] = list(self.excl_listbox.get(0, "end"))
        self.cfg["included_attachments"] = list(self.incl_listbox.get(0, "end"))
        try:
            save_config(self.cfg)
            self._log(f"Saved {CONFIG_PATH} (password excluded).", "ok")
        except Exception as exc:
            messagebox.showerror("Save failed", str(exc))

    def _log(self, msg, kind="log"):
        self.log.configure(state="normal")
        self.log.insert("end", msg + "\n", () if kind == "log" else (kind,))
        self.log.see("end")
        self.log.configure(state="disabled")

    def _drain_log(self):
        try:
            while True:
                kind, payload = self.log_q.get_nowait()
                if kind == "total":
                    self.progress.configure(maximum=max(payload, 1), value=0)
                elif kind == "progress":
                    self.progress.configure(value=payload)
                elif kind == "log_file":
                    self._last_log_file = payload
                elif kind == "done":
                    self.btn_run.configure(state="normal")
                    self.btn_stop.configure(state="disabled")
                    self._refresh_cache_info()
                    if payload:
                        log_line = (f"\n\nLog file:\n{self._last_log_file}"
                                   if getattr(self, "_last_log_file", None) else "")
                        uniq_line = ("\n(plus matched_uniq.csv / unmatched_uniq.csv)"
                                    if self.v_dedupe.get() else "")
                        messagebox.showinfo(
                            "Finished",
                            f"Matched: {payload[0]}\nUnmatched: {payload[1]}\n\n"
                            f"Files written to:\n{self.out_dir.get()}{uniq_line}{log_line}")
                else:
                    self._log(str(payload), kind)
        except queue.Empty:
            pass
        self.after(150, self._drain_log)

    def _run(self):
        csv_path = self.csv_path.get().strip()
        out_dir = self.out_dir.get().strip() or APP_DIR
        if not csv_path or not os.path.isfile(csv_path):
            messagebox.showwarning("Missing", "Pick a valid emails CSV.")
            return
        if not self.cfg["groups"]:
            messagebox.showwarning("Missing", "Add at least one group.")
            return
        os.makedirs(out_dir, exist_ok=True)

        options = self._collect_options()
        imap_cfg = self._imap_cfg()
        local_cfg = self._local_cfg()
        if options["fetch_attachments"] and not options["offline_mode"]:
            if options["email_source"] == "local":
                if not local_cfg["folder"]:
                    messagebox.showwarning("Missing", "Pick the extracted mail export folder "
                                                      "on the Run tab.")
                    return
                if not os.path.isdir(local_cfg["folder"]):
                    messagebox.showwarning("Missing", f"Folder not found:\n{local_cfg['folder']}")
                    return
            else:
                if not imap_cfg["server"] or not imap_cfg["username"]:
                    messagebox.showwarning("Missing", "IMAP server and username are required "
                                                      "for attachment search.")
                    return
                if not self.v_pass.get():
                    messagebox.showwarning("Missing", "Enter the IMAP password on the "
                                                      "IMAP Login tab.")
                    return

        self.cfg["imap"] = imap_cfg
        self.cfg["local"] = local_cfg
        self.cfg["options"] = options
        self.cfg.setdefault("recent", {})["last_csv"] = csv_path
        self.cfg["recent"]["last_out_dir"] = out_dir
        try:
            save_config(self.cfg)
        except Exception:
            pass

        groups = [Group(g["group_name"], g["keywords"]) for g in self.cfg["groups"]]
        self.cfg["excluded_attachments"] = list(self.excl_listbox.get(0, "end"))
        self.cfg["included_attachments"] = list(self.incl_listbox.get(0, "end"))
        self.processor = Processor(csv_path, out_dir, groups, options,
                                   imap_cfg, self.v_pass.get(), self.log_q,
                                   exclude_patterns=self.cfg["excluded_attachments"],
                                   include_patterns=self.cfg["included_attachments"],
                                   local_cfg=local_cfg)
        self.btn_run.configure(state="disabled")
        self.btn_stop.configure(state="normal")
        self._log("--- run started ---")
        self.processor.start()

    def _stop(self):
        if self.processor:
            self.processor.stop_flag.set()
            self._log("Stop requested ...", "warn")

    def _on_close(self):
        if self.processor and self.processor.is_alive():
            if not messagebox.askyesno("Quit", "A run is in progress. Quit anyway?"):
                return
            self.processor.stop_flag.set()
        self.destroy()


if __name__ == "__main__":
    App().mainloop()
