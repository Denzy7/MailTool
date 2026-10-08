"""Text extraction from attachments (.pdf, .docx, .doc) for keyword matching."""
from __future__ import annotations

import fnmatch
import os
import re

from mailtool.core.util import run
from mailtool.printing.tools import find_antiword, find_catdoc

try:
    from pypdf import PdfReader
    HAVE_PDF = True
except ImportError:
    try:
        from PyPDF2 import PdfReader  # type: ignore
        HAVE_PDF = True
    except ImportError:
        HAVE_PDF = False

try:
    import docx  # python-docx
    HAVE_DOCX = True
except ImportError:
    HAVE_DOCX = False

try:
    import docx2txt
    HAVE_DOCX2TXT = True
except ImportError:
    HAVE_DOCX2TXT = False

try:
    import olefile
    HAVE_OLEFILE = True
except ImportError:
    HAVE_OLEFILE = False

ALLOWED_EXTS = (".pdf", ".docx", ".doc")


def searchable(filename):
    return os.path.splitext(filename or "")[1].lower() in ALLOWED_EXTS


def extract_pdf(path):
    if not HAVE_PDF:
        raise RuntimeError("pypdf not installed")
    reader = PdfReader(path)
    if getattr(reader, "is_encrypted", False):
        try:
            reader.decrypt("")
        except Exception:
            raise RuntimeError("PDF is password-protected")
    out = []
    for page in reader.pages:
        try:
            out.append(page.extract_text() or "")
        except Exception:
            pass
    return "\n".join(out)


def extract_docx(path):
    if HAVE_DOCX:
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        for section in d.sections:
            for container in (section.header, section.footer):
                for p in container.paragraphs:
                    parts.append(p.text)
        return "\n".join(parts)
    if HAVE_DOCX2TXT:
        return docx2txt.process(path) or ""
    raise RuntimeError("python-docx / docx2txt not installed")


def extract_doc(path):
    for tool in (find_antiword(), find_catdoc()):
        if tool:
            rc, out, _ = run([tool, path], 60)
            if rc == 0 and out.strip():
                return out
    if HAVE_OLEFILE:
        try:
            if olefile.isOleFile(path):
                ole = olefile.OleFileIO(path)
                try:
                    if ole.exists("WordDocument"):
                        return _strip_binary(ole.openstream("WordDocument").read())
                finally:
                    ole.close()
        except Exception:
            pass
    with open(path, "rb") as fh:
        return _strip_binary(fh.read())


def _strip_binary(raw):
    text = raw.decode("latin-1", "ignore")
    text = re.sub(r"[^\x20-\x7E\r\n]+", " ", text)
    return re.sub(r"\s{2,}", " ", text)


def extract_text(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return extract_pdf(path)
    if ext == ".docx":
        return extract_docx(path)
    if ext == ".doc":
        return extract_doc(path)
    return ""


def matches_any_pattern(filename, patterns):
    """Plain substrings ('disclaimer') and globs ('*_logo.*'), case-insensitive."""
    name = (filename or "").lower()
    for pat in patterns or ():
        pat = (pat or "").strip().lower()
        if pat and (fnmatch.fnmatch(name, pat) or pat in name):
            return True
    return False


def filter_decision(filename, exclude, include):
    """'skip' | 'override' (excluded but forced back by include) | 'keep'"""
    ex = matches_any_pattern(filename, exclude)
    if not ex:
        return "keep"
    return "override" if matches_any_pattern(filename, include) else "skip"
